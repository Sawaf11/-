"""Tiny game-layer on top of python-pptx: pictures, shapes, groups, slide links, and hand-written OOXML animation
timelines (auto-start, click-triggered, sounds). Copied from fatha-adventure/src/pptx_kit.py and extended with
groups (parent=), left-to-right English text (rtl=False) and a deck title."""
import os, copy, re
from pptx import Presentation
from pptx.util import Emu
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, 'assets', 'img')
AUD = os.path.join(HERE, 'assets', 'audio')
IMG_DIRS, AUD_DIRS = [IMG], [AUD]  # a deck may put its own folders first
def find(dirs, *names):
    for d in dirs:
        for n in names:
            if os.path.exists(os.path.join(d, n)): return os.path.join(d, n)
    raise FileNotFoundError(names)
PX = 6350                      # EMU per design pixel (1920x1080 design canvas)
E = lambda v: int(round(v * PX))
NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
# Quranic writing: dagger alif, hamzat al-wasl, small waw / yaa, head-of-khaa sukun, madd sign, small high letters
QURANIC = re.compile('[\u0670\u0671\u06D6-\u06ED]')
DAGGER_RED = 'E03131'                  # the dagger alif is drawn red, as in the madd decks
SMALL_RED = '\u0670\u06E5\u06E6'
NOJOIN_LEFT = 'اأإآٱدذرزوؤةءى'          # letters that do not join the next letter
def place_dagger(t):
    """As in the mushaf: inside a word the dagger alif stands in the joint between its letter and the next one. After a
    letter that joins to the left a kashida (ـ) carries it, so the letters stay joined; elsewhere it sits on its letter."""
    return re.sub('([\u0621-\u064A\u0671])([\u064B-\u065F\u06E1]*)\u0670(?=[\u0653\u0654\u06D6-\u06ED]*[\u0621-\u064A\u0671])',
                  lambda m: m.group(1) + m.group(2) + ('' if m.group(1) in NOJOIN_LEFT else 'ـ') + '\u0670', t)
PREVIEW = False                # False | 'start' (initial state) | 'end' (everything visible) - static builds for LibreOffice previews

class Deck:
    def __init__(self, title='Games'):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = 12192000, 6858000
        cp = self.prs.core_properties; cp.title = title; cp.author = 'Claude'; cp.subject = 'Interactive reading games for children'
        self.layout = self.prs.slide_layouts[6]
        self.slides = {}
        self.snd_parts = {}
    def new(self, key, **kw):
        s = Scene(self, key, **kw); self.slides[key] = s; return s
    def sound_part(self, name):
        if name not in self.snd_parts:
            blob = open(find(AUD_DIRS, name + '.wav'), 'rb').read()
            self.snd_parts[name] = Part(PackURI(f'/ppt/media/snd_{name}.wav'), 'audio/wav', self.prs.part.package, blob)
        return self.snd_parts[name]
    def save(self, path):
        for s in self.slides.values(): s.finish()
        self.prs.save(path)

class Scene:
    def __init__(self, deck, key, fade=False, hidden_slide=False, adv=True, tr_snd=None):
        self.d, self.key, self.fade, self.hidden_slide, self.adv, self.tr_snd = deck, key, fade, hidden_slide, adv, tr_snd
        self.steps = []
        self.slide = deck.prs.slides.add_slide(deck.layout)
        self.auto, self.trig, self.links = [], [], []
        self.anchor = None
        self.snd_rids = {}
        self.music = None
    # ---- shapes ----
    def pic(self, img, x, y, w, h, name=None, hidden=False, link=None, descr='', rot=0, parent=None):
        if PREVIEW == 'start' and hidden: return None
        p = (parent or self.slide.shapes).add_picture(find(IMG_DIRS, img + '.jpeg', img + '.png'), Emu(E(x)), Emu(E(y)), Emu(E(w)), Emu(E(h)))
        p.name = name or img
        p._element.nvPicPr.cNvPr.set('descr', descr or (name or img).replace('_', ' '))
        if rot: p.rotation = rot
        if link: self.links.append((p, link))
        return p.shape_id if not PREVIEW or True else None
    def letter(self, text, cx, cy, d, color='16213E', ring=None, plain=False, fs=None, name=None, hidden=False, link=None):
        """Native, editable Arabic text (Arial, RTL, ar-SA) in a white badge (or plain). Matches the earlier fathah decks."""
        if PREVIEW == 'start' and hidden: return None
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        sp = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(E(cx - d / 2)), Emu(E(cy - d / 2)), Emu(E(d)), Emu(E(d)))
        sp.name = name or f'Arabic letter {text}'
        st = sp._element.find(f'{{{NS_P}}}style')
        if st is not None: sp._element.remove(st)
        if plain: sp.fill.background(); sp.line.fill.background()
        else:
            sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor(255, 255, 255)
            sp.line.color.rgb = RGBColor.from_string(ring or '3A8FD9'); sp.line.width = Pt(max(2, d / 24))
        tf = sp.text_frame; tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        p._p.get_or_add_pPr().set('rtl', '1')
        r = p.add_run(); r.text = text
        r.font.size = Pt(fs or d * 0.36); r.font.bold = True; r.font.name = 'Arial'; r.font.color.rgb = RGBColor.from_string(color)
        rPr = r._r.get_or_add_rPr(); rPr.set('lang', 'ar-SA')
        cs = etree.SubElement(rPr, f'{{{NS_A}}}cs'); cs.set('typeface', 'Arial')
        sp._element.nvSpPr.cNvPr.set('descr', f'Arabic letter {text}')
        if link: self.links.append((sp, link))
        return sp.shape_id
    def shp(self, kind, x, y, w, h, fill=None, line=None, lw=0, name='shape', text=None, fs=None, color='16213E', rect_r=None, dash=False, link=None, rot=0, alpha=None, hidden=False, bold=False, parent=None, rtl=True):
        """Native auto-shape (+ optional Arabic text: Arial, RTL, ar-SA). x,y,w,h in design px; fs in pt; lw in pt."""
        if PREVIEW == 'start' and hidden: return None
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        sp = (parent or self.slide.shapes).add_shape(getattr(MSO_SHAPE, {'STAR_5': 'STAR_5_POINT'}.get(kind, kind)), Emu(E(x)), Emu(E(y)), Emu(E(w)), Emu(E(h)))
        sp.name = name
        st = sp._element.find(f'{{{NS_P}}}style')
        if st is not None: sp._element.remove(st)
        if fill:
            sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
            if alpha is not None:
                clr = sp._element.spPr.find(f'{{{NS_A}}}solidFill')[0]; etree.SubElement(clr, f'{{{NS_A}}}alpha', val=str(int(alpha * 1000)))
        elif text is None and fill is None: sp.fill.background()
        else: sp.fill.background()
        if line: sp.line.color.rgb = RGBColor.from_string(line); sp.line.width = Pt(lw)
        else: sp.line.fill.background()
        if rect_r is not None and kind == 'ROUNDED_RECTANGLE': sp.adjustments[0] = rect_r
        if dash and line:
            ln = sp._element.spPr.find(f'{{{NS_A}}}ln'); etree.SubElement(ln, f'{{{NS_A}}}prstDash', val='dash')
        if rot: sp.rotation = rot
        if text is not None:
            tf = sp.text_frame; tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p._p.get_or_add_pPr().set('rtl', '1' if rtl else '0')
            runs = text if isinstance(text, list) else [(text, color)]
            plain = ''.join(t for t, _ in runs)
            font = 'Traditional Arabic' if QURANIC.search(plain) else 'Arial'   # Arial has no Quranic marks
            if font != 'Arial':                    # dagger alif in the joint (on a kashida); small alif / waw / yaa red
                runs = [(part, DAGGER_RED if part in SMALL_RED else col) for tx, col in runs for part in re.split('([\u0670\u06E5\u06E6])', place_dagger(tx)) if part]
            for tx, col in runs:
                r = p.add_run(); r.text = tx
                r.font.size = Pt(fs); r.font.bold = bold; r.font.name = font; r.font.color.rgb = RGBColor.from_string(col)
                rPr = r._r.get_or_add_rPr(); rPr.set('lang', 'ar-SA' if rtl else 'en-US'); etree.SubElement(rPr, f'{{{NS_A}}}cs').set('typeface', font)
            sp._element.nvSpPr.cNvPr.set('descr', ''.join(t for t, _ in runs))
        if link: self.links.append((sp, link))
        return sp.shape_id
    def hot(self, x, y, w, h, name='hotspot', link=None):
        from pptx.enum.shapes import MSO_SHAPE
        sp = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(E(x)), Emu(E(y)), Emu(E(w)), Emu(E(h)))
        sp.name = name
        spPr = sp._element.spPr
        for ch in list(spPr):
            if ch.tag.endswith('}solidFill') or ch.tag.endswith('}ln') or ch.tag.endswith('}noFill'): spPr.remove(ch)
        fill = etree.SubElement(spPr, f'{{{NS_A}}}solidFill'); clr = etree.SubElement(fill, f'{{{NS_A}}}srgbClr', val='FFFFFF')
        etree.SubElement(clr, f'{{{NS_A}}}alpha', val='1000')      # 1% opaque: invisible but click-able
        ln = etree.SubElement(spPr, f'{{{NS_A}}}ln'); etree.SubElement(ln, f'{{{NS_A}}}noFill')
        st = sp._element.find(f'{{{NS_P}}}style')
        if st is not None: sp._element.remove(st)
        if link: self.links.append((sp, link))
        return sp.shape_id
    # ---- timeline ----
    def sound_rid(self, name):
        if name not in self.snd_rids:
            self.snd_rids[name] = self.slide.part.relate_to(self.d.sound_part(name), RT.AUDIO)
        return self.snd_rids[name]
    def need_anchor(self):
        if self.anchor is None:
            p = self.slide.shapes.add_picture(os.path.join(IMG, 'blank.png'), Emu(E(2)), Emu(E(2)), Emu(E(4)), Emu(E(4)))
            p.name = 'sound anchor'; p._element.nvPicPr.cNvPr.set('descr', 'sound anchor (invisible)')
            self.anchor = p.shape_id
        return self.anchor
    def on_load(self, effects): self.auto += effects
    def on_step(self, effects): self.steps.append(effects)
    def on_click(self, spid, effects, restart='whenNotActive'): self.trig.append((spid, effects, restart))
    def add_music(self, name):
        self.music = name
    def finish(self):
        for shp, tgt in self.links:
            shp.click_action.target_slide = self.d.slides[tgt].slide if isinstance(tgt, str) else tgt
        if PREVIEW: return
        sld = self.slide._element
        if self.hidden_slide: sld.set('show', '0')
        if not self.adv or self.fade or self.tr_snd:
            tr = etree.SubElement(sld, f'{{{NS_P}}}transition')
            if not self.adv: tr.set('advClick', '0')
            if self.fade: tr.set('spd', 'med'); etree.SubElement(tr, f'{{{NS_P}}}fade')
            if self.tr_snd:                        # the sound plays whenever the slide opens, also through a link
                rid = self.sound_rid(self.tr_snd)
                s = etree.SubElement(etree.SubElement(etree.SubElement(tr, f'{{{NS_P}}}sndAc'), f'{{{NS_P}}}stSnd'), f'{{{NS_P}}}snd')
                s.set(f'{{{NS_R}}}embed', rid); s.set('name', self.tr_snd + '.wav')
        if self.music: self._music_shape()
        xml = timing_xml(self)
        if xml is not None: sld.append(etree.fromstring(xml))
    def _music_shape(self):
        pt = self.d.sound_part(self.music)
        rid_a = self.slide.part.relate_to(pt, RT.AUDIO)
        rid_m = self.slide.part.relate_to(pt, 'http://schemas.microsoft.com/office/2007/relationships/media')
        pic = self.slide.shapes.add_picture(os.path.join(IMG, 'blank.png'), Emu(E(8)), Emu(E(8)), Emu(E(40)), Emu(E(40)))
        pic.name = 'Background music (Lumi lullaby)'
        nv = pic._element.nvPicPr
        nv.cNvPr.set('descr', 'background music')
        h = etree.SubElement(nv.cNvPr, f'{{{NS_A}}}hlinkClick'); h.set(f'{{{NS_R}}}id', ''); h.set('action', 'ppaction://media')
        nvPr = nv.find(f'{{{NS_P}}}nvPr')
        af = etree.SubElement(nvPr, f'{{{NS_A}}}audioFile'); af.set(f'{{{NS_R}}}link', rid_a)
        ext_lst = etree.SubElement(nvPr, f'{{{NS_P}}}extLst'); ext = etree.SubElement(ext_lst, f'{{{NS_P}}}ext'); ext.set('uri', '{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}')
        m = etree.SubElement(ext, '{http://schemas.microsoft.com/office/powerpoint/2010/main}media', nsmap={'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main'})
        m.set(f'{{{NS_R}}}embed', rid_m)
        self.music_id = pic.shape_id

# ----------------- effects -----------------
def fx(kind, sp, t=0, **kw): return dict(kind=kind, sp=sp, t=int(t), **kw)
def snd(name, t=0, vol=100): return dict(kind='sound', sp=None, t=int(t), name=name, vol=vol)

class Ids:
    def __init__(self, n=3): self.n = n
    def __call__(self): self.n += 1; return self.n

def _tgt(spid): return f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
def _set_vis(i, spid, val, delay=0):
    return f'<p:set><p:cBhvr><p:cTn id="{i()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst></p:cTn>{_tgt(spid)}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="{val}"/></p:to></p:set>'
def _rep(e):
    s = ''
    if e.get('repeat'): s += f' repeatCount="{ "indefinite" if e["repeat"] == "inf" else int(e["repeat"] * 1000)}"'
    if e.get('rev'): s += ' autoRev="1"'
    return s

def effect_xml(e, i, scene, node):
    k, t, sp = e['kind'], e['t'], e['sp']
    pre = {'appear': (1, 'entr'), 'disappear': (1, 'exit'), 'fadein': (10, 'entr'), 'fadeout': (10, 'exit'), 'zoomin': (53, 'entr'),
           'move': (0, 'path'), 'scale': (6, 'emph'), 'spin': (8, 'emph'), 'sound': (1, 'entr')}[k]
    body = ''
    d = int(e.get('dur', 500))
    if k == 'sound':
        sp = scene.need_anchor(); rid = scene.sound_rid(e['name'])
        body += _set_vis(i, sp, 'visible')
        body += (f'<p:audio><p:cMediaNode vol="{int(e.get("vol", 100)) * 1000}"><p:cTn id="{i()}" display="0"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                 f'<p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn>'
                 f'<p:tgtEl><p:sndTgt r:embed="{rid}" name="{e["name"]}.wav"/></p:tgtEl></p:cMediaNode></p:audio>')
    elif k == 'appear': body += _set_vis(i, sp, 'visible')
    elif k == 'disappear': body += _set_vis(i, sp, 'hidden')
    elif k == 'fadein':
        body += _set_vis(i, sp, 'visible') + f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{i()}" dur="{d}"/>{_tgt(sp)}</p:cBhvr></p:animEffect>'
    elif k == 'fadeout':
        body += f'<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn id="{i()}" dur="{d}"/>{_tgt(sp)}</p:cBhvr></p:animEffect>' + _set_vis(i, sp, 'hidden', max(1, d - 1))
    elif k == 'zoomin':
        body += _set_vis(i, sp, 'visible')
        for attr in ('ppt_w', 'ppt_h'):
            body += (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{i()}" dur="{d}" fill="hold"/>{_tgt(sp)}<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
                     f'<p:tavLst><p:tav tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#{attr}"/></p:val></p:tav></p:tavLst></p:anim>')
        body += f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{i()}" dur="{d}"/>{_tgt(sp)}</p:cBhvr></p:animEffect>'
    elif k == 'move':
        pts = e.get('path') or [(e['dx'], e['dy'])]
        path = 'M 0 0 ' + ' '.join(f'L {x:.4f} {y:.4f}' for x, y in pts) + ' E'
        body += (f'<p:animMotion origin="layout" path="{path}" pathEditMode="relative" ptsTypes="{"A" * (len(pts) + 1)}"><p:cBhvr><p:cTn id="{i()}" dur="{d}" fill="hold"{_rep(e)}/>{_tgt(sp)}'
                 f'<p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>')
    elif k == 'scale':
        body += f'<p:animScale><p:cBhvr><p:cTn id="{i()}" dur="{d}" fill="hold"{_rep(e)}/>{_tgt(sp)}</p:cBhvr><p:by x="{int(e["sx"] * 1000)}" y="{int(e["sy"] * 1000)}"/></p:animScale>'
    elif k == 'spin':
        body += f'<p:animRot by="{int(e["deg"] * 60000)}"><p:cBhvr><p:cTn id="{i()}" dur="{d}" fill="hold"{_rep(e)}/>{_tgt(sp)}<p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>'
    if e.get('sound') and k != 'sound':
        rid = scene.sound_rid(e['sound'])
        body += (f'<p:audio><p:cMediaNode vol="{int(e.get("vol", 90)) * 1000}"><p:cTn id="{i()}" display="0"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                 f'<p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn>'
                 f'<p:tgtEl><p:sndTgt r:embed="{rid}" name="{e["sound"]}.wav"/></p:tgtEl></p:cMediaNode></p:audio>')
    return (f'<p:par><p:cTn id="{i()}" presetID="{pre[0]}" presetClass="{pre[1]}" presetSubtype="0" fill="hold" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="{t}"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')

def timing_xml(sc):
    if not sc.auto and not sc.trig and not sc.music and not sc.steps: return None
    i = Ids(3)
    effects = list(sc.auto)
    # music: media call on load
    music_par = ''
    if sc.music:
        mid = sc.music_id
        music_par = (f'<p:par><p:cTn id="{i()}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="withEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                     f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{i()}" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{mid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>')
    auto_inner = ''.join(effect_xml(e, i, sc, 'withEffect') for e in effects) + music_par
    main_kids = ''
    if auto_inner:
        main_kids = (f'<p:par><p:cTn id="{i()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
                     f'<p:par><p:cTn id="{i()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{auto_inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    for st in sc.steps:
        inner = ''.join(effect_xml(e, i, sc, 'clickEffect' if n == 0 else 'withEffect') for n, e in enumerate(st))
        main_kids += (f'<p:par><p:cTn id="{i()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                      f'<p:par><p:cTn id="{i()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    trig_xml = ''
    for spid, effs, restart in sc.trig:
        inner = ''.join(effect_xml(e, i, sc, 'clickEffect' if n == 0 else 'withEffect') for n, e in enumerate(sorted(effs, key=lambda z: z['t'])))
        trig_xml += (f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{i()}" restart="{restart}" fill="hold" evtFilter="cancelBubble" nodeType="interactiveSeq">'
                     f'<p:stCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cond></p:stCondLst><p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync><p:childTnLst>'
                     f'<p:par><p:cTn id="{i()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="{i()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
                     f'</p:childTnLst></p:cTn></p:seq>')
    music_node = ''
    if sc.music:
        music_node = (f'<p:audio><p:cMediaNode vol="55000" numSld="999" showWhenStopped="0"><p:cTn id="{i()}" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst>'
                      f'<p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sc.music_id}"/></p:tgtEl></p:cMediaNode></p:audio>')
    xml = (f'<p:timing xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           + (f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{main_kids}</p:childTnLst></p:cTn>'
              f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
              if main_kids else '')                # a slide with only tap triggers has no main sequence (an empty one is not valid)
           + f'{trig_xml}{music_node}</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    return xml.encode()

# ----------------- Lumi helpers -----------------
def lbox(cx, feet_y, size): return (cx - size / 2, feet_y - size * 0.96, size, size)

def pose_timeline(sc, box, tl, name_prefix='Hero', img_prefix='c_hero_'):
    """tl = [(pose, t_ms), ...]; first pose is visible at load. Returns ({pose: spid}, effects)."""
    ids, effs = {}, []
    first = tl[0][0]
    for pose in dict.fromkeys(p for p, _ in tl):
        ids[pose] = sc.pic(img_prefix + pose, *box, name=f'{name_prefix} ({pose})', hidden=(pose != first), descr=f'{name_prefix}, {pose} pose')
    for n, (pose, t) in enumerate(tl):
        if n > 0: effs.append(fx('appear', ids[pose], t))
        if n + 1 < len(tl): effs.append(fx('disappear', ids[pose], tl[n + 1][1]))
    return ids, effs

def hop(ids, t, dy=-0.07, dur=260, times=1):
    return [fx('move', sp, t, dx=0, dy=dy, dur=dur, rev=True, repeat=times if times > 1 else None) for sp in ids.values()]

def group(sc, name, hidden=False):
    """An empty group on the slide; add pictures / shapes with parent=g.shapes. Returns (group, its shape id)."""
    if PREVIEW == 'start' and hidden: return None, None
    g = sc.slide.shapes.add_group_shape(); g.name = name
    return g, g.shape_id
